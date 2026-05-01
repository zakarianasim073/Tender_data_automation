"""DOCX template generator with dynamic placeholders and repeating rows."""

from __future__ import annotations

import copy
from dataclasses import asdict
from pathlib import Path
from typing import Any

from ..models.tender_data import TenderData


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        return f"{value:g}"
    return str(value)


def _build_replacements(td: TenderData) -> dict[str, str]:
    data = asdict(td)
    replacements = {
        "{{" + key.upper() + "}}": _stringify(value)
        for key, value in data.items()
        if not isinstance(value, (list, dict))
    }
    replacements.update({
        "{{TENDER_SECURITY_WORDS}}": td.tender_security_amount_words,
        "{{TENDER_SECURITY_NUMERIC}}": td.tender_security_bdt,
    })
    return replacements


def _merge_runs_in_paragraph(paragraph) -> None:
    if len(paragraph.runs) <= 1:
        return
    full_text = "".join(run.text for run in paragraph.runs)
    paragraph.runs[0].text = full_text
    for run in paragraph.runs[1:]:
        run.text = ""


def _replace_in_paragraph(paragraph, replacements: dict[str, str]) -> None:
    _merge_runs_in_paragraph(paragraph)
    for run in paragraph.runs:
        for token, value in replacements.items():
            if token in run.text:
                run.text = run.text.replace(token, value)


def _text_nodes(row):
    return row._tr.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")


def _fill_row(row_xml, values: list[str]):
    nodes = [node for node in _text_nodes(type("Row", (), {"_tr": row_xml})()) if node.text and node.text.strip()]
    for idx, node in enumerate(nodes):
        if idx < len(values):
            node.text = _stringify(values[idx])
    return row_xml


def _equipment_rows(template_row, td: TenderData):
    rows = []
    for item in td.equipment:
        row_xml = copy.deepcopy(template_row._tr)
        rows.append(_fill_row(row_xml, [item.sl_no, item.equipment_type, item.minimum_number]))
    return rows


def _manpower_rows(template_row, td: TenderData):
    rows = []
    for item in td.manpower:
        row_xml = copy.deepcopy(template_row._tr)
        rows.append(_fill_row(row_xml, [
            item.sl_no, item.post, item.qualification,
            item.nos, item.total_exp, item.similar_exp,
        ]))
    return rows


def _jv_partner_rows(template_row, td: TenderData):
    rows = []
    for partner in td.jv_partners:
        row_xml = copy.deepcopy(template_row._tr)
        share = f"{partner.share_percent:g}%"
        if partner.share_words:
            share += f" ({partner.share_words})"
        rows.append(_fill_row(row_xml, [
            partner.name, partner.signatory_name, partner.position,
            partner.role, share, partner.address,
        ]))
    return rows


def _replace_table_markers(table, replacements: dict[str, str], td: TenderData) -> None:
    from docx.oxml.ns import qn

    markers = {
        "{{EQUIPMENT_ROWS}}": _equipment_rows,
        "{{MANPOWER_ROWS}}": _manpower_rows,
        "{{JV_PARTNER_ROWS}}": _jv_partner_rows,
    }
    table_xml = table._tbl
    original_rows = list(table.rows)

    for idx, row in reversed(list(enumerate(original_rows))):
        row_text = " ".join(cell.text for cell in row.cells)
        matched = next((marker for marker in markers if marker in row_text), None)
        if not matched:
            continue
        new_rows = markers[matched](row, td)
        ref = row._tr
        insert_at = list(table_xml).index(ref)
        table_xml.remove(ref)
        for row_xml in reversed(new_rows):
            table_xml.insert(insert_at, row_xml)

    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                _replace_in_paragraph(paragraph, replacements)


def fill_docx_template(template_path: str, output_path: str, td: TenderData) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError("python-docx is required. Run: pip install python-docx") from exc

    doc = Document(template_path)
    replacements = _build_replacements(td)

    for paragraph in doc.paragraphs:
        _replace_in_paragraph(paragraph, replacements)
    for table in doc.tables:
        _replace_table_markers(table, replacements, td)

    for section in doc.sections:
        headers = [section.header, section.first_page_header, section.even_page_header]
        footers = [section.footer, section.first_page_footer, section.even_page_footer]
        for part in headers + footers:
            for paragraph in part.paragraphs:
                _replace_in_paragraph(paragraph, replacements)
            for table in part.tables:
                _replace_table_markers(table, replacements, td)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f"  [OK] Generated: {Path(output_path).name}")
    return output_path
