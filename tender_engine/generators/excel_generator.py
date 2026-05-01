"""
Excel Generator
---------------
Generates two Excel outputs from TenderData:

1. BOQ Excel  (BOQ_541339.xlsx)
   Mirrors the layout of BOQ_541339_2.26 Less.xlsx:
   - Header block: tender ID, procuring entity, package no, work name
   - Summary row: lot no, work description, location, security, work time, quoted %
   - BOQ item rows: item no, code, description, qty, unit, bwdb rate, amount, quoted rate, amount, %
   - Totals row

2. Work Plan Excel  (Work_Plan_541339.xlsx)
   Mirrors 7. Work Plan_541339.xlsx:
   - Header: package no, tender ID, work name
   - Year/month header rows
   - Activity rows (blank Gantt cells for user to fill)
"""

from pathlib import Path
from typing import List
from ..models.tender_data import TenderData, BOQItem, WorkActivity


# ── BOQ Excel ────────────────────────────────────────────────────────────────

def generate_boq_excel(output_path: str, td: TenderData) -> str:
    """Generate BOQ comparison Excel."""
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError("openpyxl is required. Run: pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quot"

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    header_fill = PatternFill("solid", fgColor="BDD7EE")  # light blue
    yellow_fill = PatternFill("solid", fgColor="FFFF00")

    # ── Row 2: Tender ID + Procuring entity ─────────────────────────
    ws.merge_cells("B2:C2")
    _wc(ws, "B2", f"Tender ID: {td.tender_id}", bold, center, header_fill, border)
    ws.merge_cells("D2:H2")
    _wc(ws, "D2", td.procuring_entity, bold, center, header_fill, border)
    ws.merge_cells("I2:L2")
    _wc(ws, "I2", "Quoted Rate (%)", bold, center, header_fill, border)

    # ── Row 3: Departmental Estimation label ─────────────────────────
    ws.merge_cells("D3:H3")
    _wc(ws, "D3", "Departmental Estimation", bold, center, header_fill, border)

    # ── Row 5: Column headers ────────────────────────────────────────
    headers_row5 = ["", "Lot No.", "Lot No.", "Lot Description", "Location",
                    "Tender Security", "Tender Security", "Work Time",
                    "Quoted Rate (%)", "", "", ""]
    for col, val in enumerate(headers_row5, 1):
        cell = ws.cell(row=5, column=col, value=val)
        cell.font = bold
        cell.alignment = center
        cell.border = border

    # ── Row 6: Summary data row ──────────────────────────────────────
    work_period = f"{td.start_date}\nto \n{td.completion_date}\n"
    quoted_pct = td.quoted_rate_percent  # e.g. -0.02255
    row6 = ["", td.package_no, td.package_no, td.work_name, td.location,
            td.tender_security_amount, td.tender_security_amount,
            work_period, quoted_pct, quoted_pct, quoted_pct, quoted_pct]
    for col, val in enumerate(row6, 1):
        cell = ws.cell(row=6, column=col, value=val)
        cell.alignment = center
        cell.border = border

    # ── Row 8: BOQ section header ─────────────────────────────────────
    ws.merge_cells("B8:C8")
    _wc(ws, "B8", f"Tender ID: {td.tender_id}", bold, center, header_fill, border)
    ws.merge_cells("D8:H8")
    _wc(ws, "D8", f"Bill of Quantities (BOQ)_{td.rate_schedule_ref}", bold, center, header_fill, border)
    ws.merge_cells("I8:L8")
    _wc(ws, "I8", "Quoted Rate", bold, center, header_fill, border)

    # ── Row 9: BOQ column headers ─────────────────────────────────────
    boq_headers = ["", "Item No.", "Item Code", "Item Description", "Quantity", "Unit",
                   "BWDB\nRate", "BWDB\nAmount", "Quoted\nRate", "CFT Rate",
                   "Quoted\nAmount", "%"]
    for col, val in enumerate(boq_headers, 1):
        cell = ws.cell(row=9, column=col, value=val)
        cell.font = bold
        cell.alignment = center
        cell.fill = header_fill
        cell.border = border

    # ── BOQ data rows ─────────────────────────────────────────────────
    data_start = 10
    for i, item in enumerate(td.boq_items):
        r = data_start + i
        row_data = [
            "", item.item_no, item.item_code, item.description,
            item.quantity, item.unit, item.bwdb_rate, item.bwdb_amount,
            item.quoted_rate, "-", item.quoted_amount, item.percent_diff
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.border = border
            cell.alignment = left if col == 4 else center

    # ── Totals row ───────────────────────────────────────────────────
    last_r = data_start + len(td.boq_items)
    ws.cell(row=last_r, column=8, value=td.departmental_estimate).font = bold
    ws.cell(row=last_r, column=11, value=td.quoted_total).font = bold
    ws.cell(row=last_r, column=12, value=td.quoted_rate_percent).font = bold

    # ── Column widths ────────────────────────────────────────────────
    col_widths = [2, 12, 14, 55, 12, 8, 12, 16, 12, 12, 16, 10]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[6].height = 60

    wb.save(output_path)
    print(f"  [OK] Generated: {Path(output_path).name}")
    return output_path


# ── Work Plan Excel ──────────────────────────────────────────────────────────

def generate_work_plan_excel(output_path: str, td: TenderData) -> str:
    """Generate Work Plan (Gantt) Excel."""
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError("openpyxl is required. Run: pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = td.tender_id

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    header_fill = PatternFill("solid", fgColor="BDD7EE")

    # Determine month columns from td.work_months
    months = td.work_months if td.work_months else [
        "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
        "Jan", "Feb", "Mar", "Apr", "May", "Jun"
    ]
    n_months = len(months)

    # ── Row 2: Package + WORK SCHEDULE + Tender ID ───────────────────
    ws.merge_cells(f"A2:B2")
    _wc(ws, "A2", f"Package :-{td.package_no}", bold, center, header_fill, border)
    ws.merge_cells(f"C2:{get_column_letter(2 + n_months // 2)}2")
    _wc(ws, "C2", "WORK SCHEDULE", bold, center, header_fill, border)
    tender_id_col = get_column_letter(3 + n_months // 2)
    ws.merge_cells(f"{tender_id_col}2:{get_column_letter(2 + n_months)}2")
    _wc(ws, f"{tender_id_col}2", f"Tender ID: {td.tender_id}", bold, center, header_fill, border)

    # ── Row 3-4: Work name (merged across all cols) ───────────────────
    total_cols = 2 + n_months
    ws.merge_cells(f"A3:{get_column_letter(total_cols)}4")
    _wc(ws, "A3", f'Name of Work: "{td.work_name}"', bold, center, None, border)
    ws.row_dimensions[3].height = 50

    # ── Row 6: Year headers ───────────────────────────────────────────
    _wc(ws, "A6", "Sl. No.", bold, center, header_fill, border)
    _wc(ws, "B6", "Activity", bold, center, header_fill, border)
    # Determine year for each month
    start_year = td.work_start_year
    end_year = td.work_end_year
    year_for_month = _month_years(months, start_year, end_year)
    prev_year = None
    year_start_col = 3
    for m_idx, (month, year) in enumerate(zip(months, year_for_month)):
        col_letter = get_column_letter(3 + m_idx)
        ws.cell(row=6, column=3 + m_idx, value=year).font = bold
        ws.cell(row=6, column=3 + m_idx).alignment = center
        ws.cell(row=6, column=3 + m_idx).fill = header_fill
        ws.cell(row=6, column=3 + m_idx).border = border

    # ── Row 7: Month headers ──────────────────────────────────────────
    _wc(ws, "A7", "Sl. No.", bold, center, header_fill, border)
    _wc(ws, "B7", "Activity", bold, center, header_fill, border)
    for m_idx, month in enumerate(months):
        col_letter = get_column_letter(3 + m_idx)
        _wc(ws, f"{col_letter}7", month, bold, center, header_fill, border)

    # ── Activity rows ─────────────────────────────────────────────────
    activities = td.work_activities if td.work_activities else _default_activities()
    for a_idx, act in enumerate(activities):
        r = 8 + a_idx
        ws.cell(row=r, column=1, value=act.sl_no).border = border
        ws.cell(row=r, column=2, value=act.activity).alignment = left
        ws.cell(row=r, column=2).border = border
        ws.row_dimensions[r].height = 30
        for m_idx in range(n_months):
            ws.cell(row=r, column=3 + m_idx).border = border

    # ── Column widths ─────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 42
    for m_idx in range(n_months):
        ws.column_dimensions[get_column_letter(3 + m_idx)].width = 6

    wb.save(output_path)
    print(f"  [OK] Work Plan generated: {Path(output_path).name}")
    return output_path


# ── Helpers ──────────────────────────────────────────────────────────────────

def _wc(ws, cell_ref, value, font=None, alignment=None, fill=None, border=None):
    """Write cell with optional formatting."""
    c = ws[cell_ref]
    c.value = value
    if font:
        c.font = font
    if alignment:
        c.alignment = alignment
    if fill:
        c.fill = fill
    if border:
        c.border = border


def _month_years(months: List[str], start_year: int, end_year: int) -> List[int]:
    """Assign a year to each month label based on calendar order."""
    _order = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
              "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    years = []
    current_year = start_year
    prev_idx = -1
    for m in months:
        idx = next((i for i, mo in enumerate(_order) if mo.lower() == m.lower()[:3]), 0)
        if idx < prev_idx:
            current_year += 1
        years.append(current_year)
        prev_idx = idx
    return years


def _default_activities() -> List[WorkActivity]:
    from ..models.tender_data import WorkActivity
    return [
        WorkActivity("i",    "Site Preparation & Mobilization."),
        WorkActivity("ii",   "Procurement of Man power, Equipment, Materials & etc."),
        WorkActivity("iii",  "E/W in excavation/re-excavation of khal & foundation trench etc."),
        WorkActivity("iv",   "Manufacturing of C.C. Blocks"),
        WorkActivity("v",    "Supply of Sand Filter, Geo-Textile Filter, Khoa filter etc."),
        WorkActivity("vi",   "Dumping & Placing C.C Blocks"),
        WorkActivity("vii",  "Construction Of Herring Bone Bond (HBB) Brick Road."),
        WorkActivity("viii", "Supplying, Filling & Laying Geo-Tube"),
        WorkActivity("ix",   "Fine dressing and close turfing of the slopes and the crest of embankment"),
        WorkActivity("x",    "De-mobilization & Site handover"),
    ]
