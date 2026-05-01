"""
Local BOQ Cost Prediction Model
-------------------------------
Trains and runs fully on the local machine. It uses generated tender history
or an uploaded CSV to predict BOQ item rates and flag possible anomalies.
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

BASE_DIR = Path(__file__).resolve().parents[1]
PROJECT_DIR = BASE_DIR.parent
OUTPUT_DIR = BASE_DIR / "output"
MODEL_DIR = BASE_DIR / "model_store"
MODEL_PATH = MODEL_DIR / "boq_cost_predictor.joblib"
METADATA_PATH = MODEL_DIR / "boq_cost_predictor_metadata.json"


FEATURE_COLUMNS = ["category", "region", "unit_type", "month", "year"]
TARGET_COLUMN = "rate"


class BOQCostModelError(RuntimeError):
    """Raised when the local BOQ model cannot train or predict."""


def train_model_from_outputs(output_dir: str | Path = OUTPUT_DIR) -> Dict[str, Any]:
    rows = _rows_from_generated_outputs(Path(output_dir))
    if len(rows) < 5:
        raise BOQCostModelError(
            "Need at least 5 BOQ history rows to train the local model. "
            "Generate a few tender packages first or upload a historical CSV."
        )
    return _train_from_rows(rows, source="generated_outputs")


def train_model_from_csv(csv_path: str | Path) -> Dict[str, Any]:
    try:
        import pandas as pd
    except ImportError as exc:
        raise BOQCostModelError("pandas is required. Install requirements_engine.txt first.") from exc

    df = pd.read_csv(csv_path)
    required = set(FEATURE_COLUMNS + [TARGET_COLUMN])
    missing = sorted(required - set(df.columns))
    if missing:
        raise BOQCostModelError(
            "Historical CSV is missing required columns: " + ", ".join(missing)
        )
    rows = df[FEATURE_COLUMNS + [TARGET_COLUMN]].to_dict("records")
    return _train_from_rows(rows, source=str(csv_path))


def predict_tender_costs(tender_id: str, output_dir: str | Path = OUTPUT_DIR) -> Dict[str, Any]:
    model_bundle = _load_model_bundle()
    tender_json = Path(output_dir) / tender_id / "extracted_data.json"
    if not tender_json.exists():
        raise BOQCostModelError(f"extracted_data.json not found for tender {tender_id}: {tender_json}")

    with open(tender_json, "r", encoding="utf-8") as f:
        tender = json.load(f)

    boq_items = tender.get("boq_items", [])
    if not boq_items:
        raise BOQCostModelError(f"No BOQ items found in {tender_json}")

    rows = [_prediction_features(item, tender) for item in boq_items]
    predictions = _predict_rows(model_bundle, rows)

    result_rows = []
    total_predicted = 0.0
    total_existing = 0.0
    for item, prediction in zip(boq_items, predictions):
        quantity = _to_float(item.get("quantity"))
        existing_rate = _to_float(item.get("quoted_rate") or item.get("bwdb_rate"))
        predicted_rate = round(float(prediction), 3)
        predicted_amount = round(quantity * predicted_rate, 3)
        existing_amount = round(quantity * existing_rate, 3)
        variance_pct = _variance(existing_rate, predicted_rate)
        status = "ANOMALY" if abs(variance_pct) >= 15 else "NORMAL"

        total_predicted += predicted_amount
        total_existing += existing_amount
        result_rows.append({
            "item_no": item.get("item_no"),
            "item_code": item.get("item_code"),
            "description": item.get("description"),
            "unit": item.get("unit"),
            "quantity": quantity,
            "existing_rate": existing_rate,
            "predicted_rate": predicted_rate,
            "existing_amount": existing_amount,
            "predicted_amount": predicted_amount,
            "variance_percent": round(variance_pct, 2),
            "status": status,
        })

    result = {
        "tender_id": tender_id,
        "model_path": str(MODEL_PATH),
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "total_existing_amount": round(total_existing, 3),
        "total_predicted_amount": round(total_predicted, 3),
        "total_variance_percent": round(_variance(total_existing, total_predicted), 2),
        "anomaly_count": sum(1 for row in result_rows if row["status"] == "ANOMALY"),
        "rows": result_rows,
    }

    tender_out = Path(output_dir) / tender_id
    tender_out.mkdir(parents=True, exist_ok=True)
    json_path = tender_out / f"AI_Cost_Prediction-{tender_id}.json"
    xlsx_path = tender_out / f"AI_Cost_Prediction-{tender_id}.xlsx"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(result, f, indent=2, ensure_ascii=False)
    _write_prediction_excel(result, xlsx_path)
    result["json_path"] = str(json_path)
    result["excel_path"] = str(xlsx_path)
    return result


def model_status() -> Dict[str, Any]:
    if not MODEL_PATH.exists():
        return {"trained": False, "message": "No local BOQ cost model trained yet."}
    metadata = {}
    if METADATA_PATH.exists():
        with open(METADATA_PATH, "r", encoding="utf-8") as f:
            metadata = json.load(f)
    return {"trained": True, "model_path": str(MODEL_PATH), **metadata}


def _train_from_rows(rows: List[Dict[str, Any]], source: str) -> Dict[str, Any]:
    try:
        import joblib
        import pandas as pd
        from sklearn.ensemble import GradientBoostingRegressor
        from sklearn.preprocessing import LabelEncoder
    except ImportError as exc:
        raise BOQCostModelError(
            "Missing ML dependencies. Install requirements_engine.txt first."
        ) from exc

    df = pd.DataFrame(rows)
    df = df.dropna(subset=[TARGET_COLUMN])
    for col in FEATURE_COLUMNS:
        df[col] = df[col].fillna("unknown").astype(str)
    df[TARGET_COLUMN] = df[TARGET_COLUMN].astype(float)

    if len(df) < 5:
        raise BOQCostModelError("Need at least 5 usable rows after cleaning to train.")

    encoders = {}
    x = df[FEATURE_COLUMNS].copy()
    for col in FEATURE_COLUMNS:
        enc = LabelEncoder()
        x[col] = enc.fit_transform(x[col])
        encoders[col] = enc

    model = GradientBoostingRegressor(n_estimators=100, random_state=42)
    model.fit(x, df[TARGET_COLUMN])

    MODEL_DIR.mkdir(parents=True, exist_ok=True)
    import joblib
    joblib.dump({"model": model, "encoders": encoders, "features": FEATURE_COLUMNS}, MODEL_PATH)

    metadata = {
        "trained": True,
        "trained_at": datetime.now().isoformat(timespec="seconds"),
        "source": source,
        "row_count": int(len(df)),
        "features": FEATURE_COLUMNS,
        "target": TARGET_COLUMN,
    }
    with open(METADATA_PATH, "w", encoding="utf-8") as f:
        json.dump(metadata, f, indent=2)
    return {"model_path": str(MODEL_PATH), **metadata}


def _load_model_bundle() -> Dict[str, Any]:
    if not MODEL_PATH.exists():
        raise BOQCostModelError("No trained model found. Train from history or CSV first.")
    try:
        import joblib
    except ImportError as exc:
        raise BOQCostModelError("joblib is required. Install requirements_engine.txt first.") from exc
    return joblib.load(MODEL_PATH)


def _predict_rows(bundle: Dict[str, Any], rows: List[Dict[str, Any]]) -> List[float]:
    try:
        import pandas as pd
    except ImportError as exc:
        raise BOQCostModelError("pandas is required. Install requirements_engine.txt first.") from exc

    df = pd.DataFrame(rows)
    for col in FEATURE_COLUMNS:
        df[col] = df[col].fillna("unknown").astype(str)
        enc = bundle["encoders"][col]
        known = set(enc.classes_)
        df[col] = df[col].map(lambda value: value if value in known else enc.classes_[0])
        df[col] = enc.transform(df[col])
    return bundle["model"].predict(df[FEATURE_COLUMNS]).tolist()


def _rows_from_generated_outputs(output_dir: Path) -> List[Dict[str, Any]]:
    rows = []
    for json_path in output_dir.glob("*/extracted_data.json"):
        with open(json_path, "r", encoding="utf-8") as f:
            tender = json.load(f)
        for item in tender.get("boq_items", []):
            rate = _to_float(item.get("quoted_rate") or item.get("bwdb_rate"))
            if rate <= 0:
                continue
            row = _prediction_features(item, tender)
            row[TARGET_COLUMN] = rate
            rows.append(row)
    return rows


def _prediction_features(item: Dict[str, Any], tender: Dict[str, Any]) -> Dict[str, Any]:
    date_value = tender.get("closing_date") or tender.get("publication_date") or "01-Jan-2024"
    month, year = _month_year(date_value)
    return {
        "category": _infer_category(item.get("item_code"), item.get("description")),
        "region": _infer_region(tender.get("location")),
        "unit_type": str(item.get("unit") or "unknown").strip().lower(),
        "month": str(month),
        "year": str(year),
    }


def _infer_category(item_code: Optional[str], description: Optional[str]) -> str:
    code = str(item_code or "").strip()
    text = str(description or "").lower()
    prefix = code.split("-")[0].strip() if code else ""
    if prefix:
        return prefix
    if "earth" in text:
        return "earthwork"
    if "block" in text or "concrete" in text:
        return "concrete"
    if "geo" in text:
        return "geotextile"
    if "brick" in text or "hbb" in text:
        return "road"
    return "general"


def _infer_region(location: Optional[str]) -> str:
    value = str(location or "unknown").strip().lower()
    if "," in value:
        return value.split(",")[-1].strip()
    return value or "unknown"


def _month_year(value: str) -> tuple[int, int]:
    for fmt in ("%d-%b-%Y", "%d-%B-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            dt = datetime.strptime(str(value).strip(), fmt)
            return dt.month, dt.year
        except ValueError:
            continue
    return datetime.now().month, datetime.now().year


def _to_float(value: Any) -> float:
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return 0.0


def _variance(existing: float, predicted: float) -> float:
    if existing == 0:
        return 0.0
    return ((predicted - existing) / existing) * 100


def _write_prediction_excel(result: Dict[str, Any], path: Path) -> None:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AI Prediction"
    headers = [
        "Item No", "Item Code", "Description", "Unit", "Quantity",
        "Existing Rate", "Predicted Rate", "Existing Amount", "Predicted Amount",
        "Variance %", "Status"
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="BDD7EE")

    for row in result["rows"]:
        ws.append([
            row["item_no"], row["item_code"], row["description"], row["unit"], row["quantity"],
            row["existing_rate"], row["predicted_rate"], row["existing_amount"],
            row["predicted_amount"], row["variance_percent"], row["status"],
        ])
    ws.append([])
    ws.append(["Total Existing", result["total_existing_amount"]])
    ws.append(["Total Predicted", result["total_predicted_amount"]])
    ws.append(["Total Variance %", result["total_variance_percent"]])
    ws.append(["Anomaly Count", result["anomaly_count"]])

    ws.column_dimensions["C"].width = 70
    for col in ["A", "B", "D", "E", "F", "G", "H", "I", "J", "K"]:
        ws.column_dimensions[col].width = 16
    wb.save(path)
