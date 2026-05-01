"""Description-based SOR suggestions for BOQ rows without usable item codes."""

from __future__ import annotations

import difflib
import json
import re
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any, Iterable, List

from ..models import BOQItem
from ..sor import active_sor_paths, detect_bwdb_zone, parse_bwdb_sor, parse_lged_sor

BASE_DIR = Path(__file__).resolve().parents[1]
OUTPUT_DIR = BASE_DIR / "output"
STOPWORDS = {
    "the", "and", "for", "with", "including", "in", "of", "to", "by", "as",
    "per", "all", "complete", "direction", "engineer", "charge", "etc", "work",
}


@dataclass
class SORSuggestion:
    item_no: int
    boq_code: str
    boq_description: str
    boq_unit: str
    quantity: float
    suggested_code: str
    suggested_description: str
    suggested_unit: str
    sor_rate: float
    confidence: float
    status: str


def suggest_for_tender(tender_id: str, source: str = "BWDB", limit: int = 3) -> dict:
    """Load extracted_data.json and produce SOR suggestions for each BOQ item."""
    folder = OUTPUT_DIR / tender_id
    data_path = folder / "extracted_data.json"
    if not data_path.exists():
        raise FileNotFoundError(f"Generate tender first; missing {data_path}")

    data = json.loads(data_path.read_text(encoding="utf-8"))
    zone = detect_bwdb_zone(data.get("location", ""))
    sor_items = _load_sor_items(source)
    rows = suggest_for_items(data.get("boq_items", []), sor_items, zone, limit)
    out_xlsx = folder / f"SOR_Suggestions-{tender_id}.xlsx"
    _write_suggestions_excel(rows, out_xlsx)
    out_json = folder / f"SOR_Suggestions-{tender_id}.json"
    payload = {
        "tender_id": tender_id,
        "source": source.upper(),
        "zone": zone,
        "suggestion_count": len(rows),
        "excel_path": str(out_xlsx),
        "rows": [asdict(row) for row in rows],
    }
    out_json.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return payload


def suggest_for_items(boq_items: Iterable[Any], sor_items: Iterable[Any], zone: str, limit: int = 3) -> List[SORSuggestion]:
    suggestions: List[SORSuggestion] = []
    searchable = [(_normalise(item.description), item) for item in sor_items]
    for raw in boq_items:
        boq = _boq_from_any(raw)
        ranked = sorted(
            ((_score(boq, norm_desc, sor), sor) for norm_desc, sor in searchable),
            key=lambda pair: pair[0],
            reverse=True,
        )[:limit]
        for score, sor in ranked:
            if score <= 0:
                continue
            status = "strong" if score >= 0.72 else "review" if score >= 0.52 else "weak"
            suggestions.append(SORSuggestion(
                item_no=boq.item_no,
                boq_code=boq.item_code,
                boq_description=boq.description,
                boq_unit=boq.unit,
                quantity=boq.quantity,
                suggested_code=sor.item_code,
                suggested_description=sor.description,
                suggested_unit=sor.unit,
                sor_rate=sor.get_rate(zone),
                confidence=round(score, 3),
                status=status,
            ))
    return suggestions


def _load_sor_items(source: str):
    paths = active_sor_paths()
    if source.upper() == "LGED":
        return parse_lged_sor(paths.get("LGED", ""))
    return parse_bwdb_sor(paths.get("BWDB", ""))


def _score(boq: BOQItem, sor_norm_desc: str, sor: Any) -> float:
    boq_norm = _normalise(boq.description)
    if not boq_norm or not sor_norm_desc:
        return 0.0
    seq = difflib.SequenceMatcher(None, boq_norm, sor_norm_desc).ratio()
    boq_tokens = set(boq_norm.split())
    sor_tokens = set(sor_norm_desc.split())
    overlap = len(boq_tokens & sor_tokens) / max(1, len(boq_tokens | sor_tokens))
    unit_bonus = 0.08 if _unit_key(boq.unit) and _unit_key(boq.unit) == _unit_key(sor.unit) else 0
    code_bonus = 0.04 if boq.item_code and boq.item_code in sor.item_code else 0
    return min(1.0, (seq * 0.55) + (overlap * 0.45) + unit_bonus + code_bonus)


def _normalise(text: str) -> str:
    text = re.sub(r"[^a-zA-Z0-9]+", " ", str(text).lower())
    tokens = [t for t in text.split() if len(t) > 2 and t not in STOPWORDS]
    return " ".join(tokens)


def _unit_key(unit: str) -> str:
    unit = str(unit or "").lower().replace(".", "").strip()
    return {"nos": "no", "nos.": "no", "each": "no", "sqm": "sqm", "cum": "cum"}.get(unit, unit)


def _boq_from_any(raw: Any) -> BOQItem:
    if isinstance(raw, BOQItem):
        return raw
    return BOQItem(
        item_no=int(raw.get("item_no", 0) or 0),
        item_code=str(raw.get("item_code", "") or ""),
        description=str(raw.get("description", "") or ""),
        quantity=float(raw.get("quantity", 0) or 0),
        unit=str(raw.get("unit", "") or ""),
        bwdb_rate=float(raw.get("bwdb_rate", 0) or 0),
        bwdb_amount=float(raw.get("bwdb_amount", 0) or 0),
        quoted_rate=float(raw.get("quoted_rate", 0) or 0),
        quoted_amount=float(raw.get("quoted_amount", 0) or 0),
        percent_diff=float(raw.get("percent_diff", 0) or 0),
    )


def _write_suggestions_excel(rows: List[SORSuggestion], output_path: Path) -> None:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SOR Suggestions"
    headers = [
        "Item", "BOQ Code", "BOQ Description", "Unit", "Qty", "Suggested SOR",
        "SOR Description", "SOR Unit", "SOR Rate", "Confidence", "Status"
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append([
            row.item_no, row.boq_code, row.boq_description, row.boq_unit,
            row.quantity, row.suggested_code, row.suggested_description,
            row.suggested_unit, row.sor_rate, row.confidence, row.status,
        ])
    widths = [8, 14, 60, 10, 14, 16, 70, 10, 14, 12, 12]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = width
    wb.save(output_path)
